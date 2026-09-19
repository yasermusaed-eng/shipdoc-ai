"""
pipeline.py — End-to-end Shipping Document Verification Pipeline.

Wires together:
1. Classification (LLM classifier with deterministic heuristic fallback)
2. Document Extraction (SI & BL field extraction)
3. Deterministic Comparison (case & whitespace normalized, numeric-aware)
4. Escalation Logic (Human-in-the-loop triggers for low confidence, missing docs, unreadable text, or high null counts)
"""

import json
import os
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

# Setup module resolution
PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "data"
SRC_DIR = PROJECT_ROOT / "src"
OUTPUT_DIR = PROJECT_ROOT / "output"

sys.path.insert(0, str(DATA_DIR))
sys.path.insert(0, str(SRC_DIR))

try:
    from loader import Inbox
except ImportError:
    from data.loader import Inbox

from classifier import classify_email
from llm_classifier import classify_email_llm
from extractor import extract_fields
from comparator import compare_documents

# Thresholds for escalation
CONFIDENCE_THRESHOLD = 0.70
MAX_ALLOWED_NULL_FIELDS = 2


def _find_si_and_bl(attachments: List[str]) -> Tuple[Optional[str], Optional[str]]:
    """Identifies the SI and BL attachment paths from an attachment list."""
    si_file = None
    bl_file = None

    for att in attachments:
        att_lower = att.lower()
        if "_si." in att_lower or ("si" in att_lower and not "bl" in att_lower):
            if not si_file:
                si_file = att
        elif "_bl." in att_lower or ("bl" in att_lower and not "si" in att_lower):
            if not bl_file:
                bl_file = att

    return si_file, bl_file


def _read_doc_text(inbox: Any, att_path: str) -> Tuple[Optional[str], Optional[str]]:
    """
    Safely reads document text from an attachment path.
    Returns (text_content, error_message).
    """
    try:
        # Check file extension
        ext = Path(att_path).suffix.lower()

        if ext == ".txt":
            content = inbox.read_text(att_path)
            if not content or not content.strip():
                return None, "empty_document"
            return content, None

        # For binary files (PDF/DOCX/XLSX), attempt standard reading or mark unreadable
        raw_bytes = inbox.read_bytes(att_path)

        if ext == ".pdf":
            try:
                import pypdf
                import io
                reader = pypdf.PdfReader(io.BytesIO(raw_bytes))
                text = "\n".join(page.extract_text() or "" for page in reader.pages)
                if text.strip():
                    return text, None
            except Exception:
                pass
            return None, "unreadable: PDF document requires OCR or PDF parsing library"

        if ext == ".docx":
            try:
                import docx
                import io
                doc = docx.Document(io.BytesIO(raw_bytes))
                text = "\n".join(p.text for p in doc.paragraphs)
                if text.strip():
                    return text, None
            except Exception:
                pass
            return None, "unreadable: DOCX document requires docx library"

        if ext in (".xlsx", ".xls"):
            try:
                import openpyxl
                import io
                wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
                lines = []
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        row_vals = [str(c) for c in row if c is not None]
                        if row_vals:
                            lines.append(" : ".join(row_vals))
                text = "\n".join(lines)
                if text.strip():
                    return text, None
            except Exception:
                pass
            return None, "unreadable: Spreadsheet format requires openpyxl"

        # Unknown binary format
        return None, f"unreadable: unsupported file format '{ext}'"

    except Exception as e:
        return None, f"unreadable: failed to read attachment ({e})"


def process_email(email: Dict[str, Any], inbox: Any) -> Dict[str, Any]:
    """
    Processes an incoming email end-to-end through classification, extraction,
    comparison, and escalation triage.

    Args:
        email: Raw email dictionary from inbox
        inbox: Inbox loader instance

    Returns:
        Dict shaped as:
        {
          "email_id": ...,
          "category": ...,
          "mismatch_found": bool | None,
          "mismatches": [...] | None,
          "escalate": bool | None,
          "escalation_reason": str | None
        }
    """
    email_id = email.get("email_id", "unknown")
    attachments = email.get("attachments", []) or []

    # 1. Classification (Try LLM; fallback to heuristic baseline if no API key or network error)
    category = "general"
    confidence = 1.0
    reasoning = ""

    use_llm = bool(os.getenv("GEMINI_API_KEY") or os.getenv("OPENAI_API_KEY"))

    if use_llm:
        try:
            llm_result = classify_email_llm(email, provider="auto")
            category = llm_result.get("category", "general")
            confidence = float(llm_result.get("confidence", 0.9))
            reasoning = llm_result.get("reasoning", "")
        except Exception:
            # Graceful fallback to deterministic classifier
            category = classify_email(email)
            confidence = 0.85
    else:
        category = classify_email(email)
        confidence = 0.85

    # Check for low classifier confidence escalation
    if confidence < CONFIDENCE_THRESHOLD:
        return {
            "email_id": email_id,
            "category": category,
            "mismatch_found": None,
            "mismatches": None,
            "escalate": True,
            "escalation_reason": f"low_confidence: classifier confidence {confidence:.2f} is below threshold {CONFIDENCE_THRESHOLD}",
        }

    # 2. If NOT document_comparison -> return classification result immediately
    # Both mismatch_found and escalate are None (null) because no document extraction or comparison was attempted
    if category != "document_comparison":
        return {
            "email_id": email_id,
            "category": category,
            "mismatch_found": None,
            "mismatches": None,
            "escalate": None,
            "escalation_reason": None,
        }

    # 3. Handle document_comparison emails
    si_path, bl_path = _find_si_and_bl(attachments)

    # Check missing attachments
    if not si_path or not bl_path:
        missing_parts = []
        if not si_path:
            missing_parts.append("SI")
        if not bl_path:
            missing_parts.append("BL")
        return {
            "email_id": email_id,
            "category": category,
            "mismatch_found": None,
            "mismatches": None,
            "escalate": True,
            "escalation_reason": f"missing_attachment: missing {' and '.join(missing_parts)} attachment",
        }

    # Read attachment text contents
    si_text, si_err = _read_doc_text(inbox, si_path)
    bl_text, bl_err = _read_doc_text(inbox, bl_path)

    if si_err or bl_err:
        reasons = []
        if si_err:
            reasons.append(f"SI ({si_path}): {si_err}")
        if bl_err:
            reasons.append(f"BL ({bl_path}): {bl_err}")
        return {
            "email_id": email_id,
            "category": category,
            "mismatch_found": None,
            "mismatches": None,
            "escalate": True,
            "escalation_reason": "unreadable: " + "; ".join(reasons),
        }

    # 4. Extract fields from both documents
    si_fields = extract_fields(si_text)
    bl_fields = extract_fields(bl_text)

    # Count null fields in extraction
    si_nulls = sum(1 for v in si_fields.values() if v is None)
    bl_nulls = sum(1 for v in bl_fields.values() if v is None)

    if si_nulls > MAX_ALLOWED_NULL_FIELDS or bl_nulls > MAX_ALLOWED_NULL_FIELDS:
        return {
            "email_id": email_id,
            "category": category,
            "mismatch_found": None,
            "mismatches": None,
            "escalate": True,
            "escalation_reason": (
                f"missing_value: extraction returned excessive null fields "
                f"(SI nulls: {si_nulls}/7, BL nulls: {bl_nulls}/7, max allowed: {MAX_ALLOWED_NULL_FIELDS})"
            ),
        }

    # 5. Compare documents deterministically
    comp_result = compare_documents(si_fields, bl_fields)

    return {
        "email_id": email_id,
        "category": category,
        "mismatch_found": comp_result["mismatch_found"],
        "mismatches": comp_result["mismatches"],
        "escalate": False,
        "escalation_reason": None,
    }


def run_pipeline_dataset(dataset_dir: Path, output_file: Path) -> List[Dict[str, Any]]:
    """Runs the full pipeline over all emails in dataset_dir and saves results."""
    inbox = Inbox(str(dataset_dir))
    emails = inbox.emails()
    total = len(emails)

    print(f"Loaded {total} emails from {dataset_dir / 'inbox'}")
    print("Processing pipeline...")

    results = []
    category_counts: Dict[str, int] = {}
    escalation_count = 0
    mismatch_count = 0
    ok_count = 0

    for i, email in enumerate(emails, 1):
        res = process_email(email, inbox)
        results.append(res)

        cat = res["category"]
        category_counts[cat] = category_counts.get(cat, 0) + 1

        if res["escalate"]:
            escalation_count += 1
        elif res["mismatch_found"] is True:
            mismatch_count += 1
        elif res["mismatch_found"] is False:
            ok_count += 1

        if i % 50 == 0 or i == total:
            print(f"  Processed {i}/{total} emails...")

    # Ensure output directory exists
    output_file.parent.mkdir(parents=True, exist_ok=True)
    with open(output_file, "w", encoding="utf-8") as f:
        json.dump(results, f, indent=2)

    print("\n" + "=" * 70)
    print("PIPELINE EXECUTION SUMMARY")
    print("=" * 70)
    print(f"Total Emails Processed : {total}")
    print(f"Results Saved To       : {output_file}")
    print("\nCategory Distribution:")
    for cat, cnt in sorted(category_counts.items(), key=lambda x: -x[1]):
        pct = (cnt / total * 100) if total else 0
        print(f"  - {cat:<22}: {cnt:>4} ({pct:>5.1f}%)")

    print("\nDocument Verification Outcomes:")
    print(f"  - No Mismatch (OK)   : {ok_count}")
    print(f"  - Mismatches Flagged : {mismatch_count}")
    print(f"  - Escalated to Review: {escalation_count}")
    print("=" * 70)

    return results


if __name__ == "__main__":
    out_path = OUTPUT_DIR / "results.json"
    run_pipeline_dataset(DATA_DIR, out_path)
